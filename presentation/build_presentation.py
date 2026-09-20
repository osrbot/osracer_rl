#!/usr/bin/env python3
"""Build the evidence-led, editable OSRACER live briefing deck."""
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "presentation" / "build" / "osracer_live_briefing.pptx"
W, H = 13.333, 7.5
INK = RGBColor(25, 30, 38)
NAVY = RGBColor(17, 43, 70)
BLUE = RGBColor(36, 120, 190)
TEAL = RGBColor(0, 153, 142)
ORANGE = RGBColor(217, 90, 52)
RED = RGBColor(190, 52, 58)
PAPER = RGBColor(250, 249, 246)
MUTED = RGBColor(93, 105, 117)
FONT = "Noto Sans CJK SC"

def box(slide, x, y, w, h, fill=PAPER, line=None, radius=False):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    shp.fill.solid(); shp.fill.fore_color.rgb = fill
    shp.line.color.rgb = fill if line is None else line
    return shp

def text(slide, value, x, y, w, h, size=18, color=INK, bold=False, align=PP_ALIGN.LEFT):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.clear(); tf.word_wrap = True; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = value; r.font.name = FONT; r.font.size = Pt(size); r.font.bold = bold; r.font.color.rgb = color
    return tb

def title(slide, section, headline, subtitle=""):
    text(slide, section.upper(), .55, .30, 3.0, .25, 9, BLUE, True)
    text(slide, headline, .55, .62, 12.0, .65, 27, INK, True)
    if subtitle: text(slide, subtitle, .57, 1.30, 11.9, .35, 11, MUTED)
    box(slide, .55, 1.72, 12.2, .025, BLUE)

def footer(slide, n):
    text(slide, f"OSRACER · 证据化直播材料  |  {n:02d}", .55, 7.12, 5.0, .18, 8, MUTED)

def image(slide, rel, x, y, w, h):
    path = ROOT / rel
    if path.exists(): slide.shapes.add_picture(str(path), Inches(x), Inches(y), Inches(w), Inches(h))
    else: box(slide, x, y, w, h, RGBColor(230, 232, 235), MUTED); text(slide, f"缺少媒体\n{rel}", x+.15, y+.15, w-.3, h-.3, 12, MUTED)

def metric(slide, number, label, x, y, color=BLUE):
    text(slide, number, x, y, 2.4, .55, 33, color, True)
    text(slide, label, x, y+.58, 2.5, .33, 11, MUTED)

def bullet(slide, value, x, y, w, color=INK):
    text(slide, "•", x, y, .22, .28, 17, BLUE, True)
    text(slide, value, x+.26, y-.02, w-.26, .45, 14, color)

def bar(slide, label, passed, total, x, y, color):
    text(slide, label, x, y, 2.4, .3, 13, INK, True)
    box(slide, x+2.45, y+.02, 5.5, .22, RGBColor(226, 231, 236))
    box(slide, x+2.45, y+.02, 5.5*passed/total, .22, color)
    text(slide, f"{passed}/{total}", x+8.1, y-.03, .8, .3, 14, color, True, PP_ALIGN.RIGHT)

def build():
    prs = Presentation(); prs.slide_width = Inches(W); prs.slide_height = Inches(H)
    blank = prs.slide_layouts[6]
    def new():
        s = prs.slides.add_slide(blank); bg = s.background.fill; bg.solid(); bg.fore_color.rgb = PAPER; return s

    s = new(); image(s, "output/racing/isaac/bahrain_c03v10c_dev24_seed0.png", 7.15, 0, 6.18, 7.5)
    box(s, 0, 0, 7.35, 7.5, PAPER); text(s, "OSRACER", .65, .72, 2.5, .35, 12, BLUE, True)
    text(s, "双引擎自主竞速\n哪些已经被证实？", .65, 1.25, 6.1, 1.35, 34, INK, True)
    text(s, "OpenUSD/MuJoCo 资产可用性 · Ackermann 超车/圈速 · 180° 回头弯漂移边界", .68, 2.82, 5.9, .55, 15, MUTED)
    box(s, .68, 4.2, 5.65, 1.25, RGBColor(232, 240, 247), radius=True)
    text(s, "原则：用原生录像、轨迹审计和失败对照说话\n不把局部样例扩展成全条件结论", .95, 4.45, 5.1, .68, 17, NAVY, True)
    footer(s, 1)

    s = new(); title(s, "EXECUTIVE SUMMARY", "两个问题，三层结论")
    for x, head, body, col in [(0.7, "资产层", "OpenUSD / Isaac Sim 与 MuJoCo 已完成加载、渲染、短程步进与可解码录像。", TEAL), (4.55, "竞速层", "双引擎已跑出整圈、有效超车和速度榜；成绩按引擎、赛道、速度、种子分开统计。", BLUE), (8.4, "实车等效层", "当前单电机四驱动作空间缺少后轮增速；现有漂移机制不能直接迁移。", ORANGE)]:
        box(s, x, 2.15, 3.45, 3.25, RGBColor(255,255,255), col, True); text(s, head, x+.25, 2.48, 2.9, .4, 22, col, True); text(s, body, x+.25, 3.1, 2.9, 1.35, 16, INK); text(s, "已证实范围 / 明确边界", x+.25, 4.75, 2.8, .25, 10, MUTED)
    text(s, "结论不是“都完成了”：资产可用性已完成；竞速能力部分完成；实车 180° 漂移尚未成立。", .7, 5.92, 11.5, .4, 19, RED, True); footer(s, 2)

    s = new(); title(s, "ASSET USABILITY", "导出的 OpenUSD 与 MuJoCo 资产：可用，但验收范围很窄")
    image(s, "output/usability/openusd/preview.png", .7, 2.05, 5.65, 3.18); image(s, "output/usability/mujoco/preview_0270.png", 6.95, 2.05, 5.65, 3.18)
    text(s, "OpenUSD / Isaac Sim 6.0.1", .7, 5.36, 5.4, .25, 16, NAVY, True); text(s, "20 网格、10 刚体、9 物理关节、6 可动自由度；60 步 PhysX 状态有限。", .7, 5.66, 5.4, .5, 13, INK)
    text(s, "MuJoCo 3.10.0", 6.95, 5.36, 5.4, .25, 16, NAVY, True); text(s, "robot.xml / scene.xml 各 500 步、无数值警告；固定基座和执行器跟踪限制仍在。", 6.95, 5.66, 5.4, .5, 13, INK)
    text(s, "不能从这两段可用性视频推断可靠驾驶、物理保真或控制器质量。", .7, 6.42, 11.6, .3, 15, RED, True); footer(s, 3)

    s = new(); title(s, "EXPERIMENT CONTRACT", "我们用受限的主动观测，而不是全局定位来闭环竞速")
    for x, label, desc in [(0.85,"15 Hz laser","laser 坐标系\n270° / 361 束"),(4.42,"轮速 + 转角","驱动状态与 Ackermann\n转向几何"),(8.0,"策略 + 安全层","预测、超车、漂移\n碰撞与出界许可")]:
        box(s,x,2.15,3.05,1.85,RGBColor(235,243,249),BLUE,True); text(s,label,x+.2,2.45,2.65,.32,19,NAVY,True); text(s,desc,x+.2,2.9,2.65,.62,15,INK,False,PP_ALIGN.CENTER)
    for x in (3.95,7.52): text(s,"→",x,2.72,.35,.4,25,BLUE,True,PP_ALIGN.CENTER)
    bullet(s,"双车闭环：有效整圈、完整超越、碰撞、出界、连续侧滑均进入同一验证口径。", .9,4.65,11.4)
    bullet(s,"不使用全局位置；录像、轨迹和独立审计必须相互对应。", .9,5.25,11.4); footer(s, 4)

    s = new(); title(s, "EVALUATION", "通过不是跑起来：必须同时满足 5 个可审计条件")
    items=[("有效整圈","完整回到终点且无判定失败"), ("有效超车","完整超越，不把名次抖动算进去"), ("连续漂移","侧滑角与持续时间达标"), ("轨迹审计","独立校验每回合状态与统计"), ("原生录像","同一标签下可解码、可回放")]
    for i,(a,b) in enumerate(items):
        x=.7+(i%3)*4.1; y=2.15+(i//3)*1.75; box(s,x,y,3.65,1.35,RGBColor(255,255,255),BLUE,True); text(s,str(i+1).zfill(2),x+.22,y+.22,.45,.3,14,BLUE,True); text(s,a,x+.78,y+.18,2.6,.3,18,INK,True); text(s,b,x+.25,y+.68,3.1,.35,12,MUTED)
    text(s,"“加载模型”“单个成功视频”或“测试通过”都不能替代整圈资格。", .7,6.05,11.3,.35,18,RED,True); footer(s,5)

    s = new(); title(s, "RACING RESULTS", "冻结 v10c：MuJoCo 24/24，Isaac 22/24；差异被保留")
    bar(s,"MuJoCo · 7.25 m/s · 24 赛道",24,24,.85,2.35,TEAL); bar(s,"Isaac · 7.25 m/s · 24 赛道",22,24,.85,3.10,BLUE)
    metric(s,"174","已审计回合",.9,4.15,BLUE); metric(s,"141","任务成功",3.75,4.15,TEAL); metric(s,"48","同批双引擎录像",6.6,4.15,ORANGE)
    box(s,9.2,2.25,3.1,3.45,RGBColor(255,244,240),ORANGE,True); text(s,"Isaac 的两条失败",9.45,2.58,2.5,.3,17,ORANGE,True); bullet(s,"Spa：起步车车接触",9.42,3.15,2.5); bullet(s,"Suzuka：桥面段车身失稳",9.42,3.82,2.5); text(s,"所以不能称 Isaac 全赛道完成。",9.45,4.75,2.35,.45,13,RED,True); footer(s,6)

    s = new(); title(s, "LAP TIME", "9.0 m/s 刷新最快圈速，但 Isaac 的稳定性付出了代价")
    bar(s,"MuJoCo · 9.0 m/s",24,24,.85,2.25,TEAL); bar(s,"Isaac · 9.0 m/s",20,24,.85,2.95,BLUE)
    metric(s,"8.97","Mu 峰值 m/s",.85,4.05,TEAL); metric(s,"9.00–9.04","Isaac 峰值 m/s",4.0,4.05,BLUE); metric(s,"1.5–10.4%","Mu 单圈加速",8.3,4.05,ORANGE)
    text(s,"MuJoCo：24 条赛道均快于 7.25 m/s 基线。Isaac：20 条可比赛道均更快，但新增 Austin、Las Vegas 停滞，Spa/Suzuka 仍失败。", .85,5.55,11.35,.6,16,INK)
    text(s,"工程结论：共享冻结配置维持 7.25 m/s；9.0 m/s 是 MuJoCo 速度榜，不是跨引擎稳定配置。", .85,6.25,11.35,.35,16,RED,True); footer(s,7)

    s = new(); title(s, "OVERTAKE", "主动超车已计入整圈资格，但只在明确的对手条件下成立")
    image(s,"output/racing/mujoco/bahrain_c03v10c_dev24_seed0.png",.7,2.05,5.9,3.45); image(s,"output/racing/isaac/bahrain_c03v10c_dev24_seed0.png",6.77,2.05,5.85,3.45)
    text(s,"MuJoCo Bahrain：代表样例",.75,5.62,4.0,.25,14,TEAL,True); text(s,"Isaac Bahrain：代表样例",6.8,5.62,4.0,.25,14,BLUE,True)
    text(s,"默认对手：同类传感器策略、巡航 2.8 m/s、发车领先 3 m。每条有效圈的“1 次”是完整超越，不代表对高速对手或任意初始位置的能力。",.75,6.12,11.65,.5,14,INK); footer(s,8)

    s = new(); title(s, "DRIFT", "180° 回头弯：仿真有漂移样例，但当前实车等效动作空间不支持它")
    image(s,"output/hairpin_fast/mujoco_fast_peak_slip.png",.75,2.0,5.75,3.55)
    metric(s,"27.0°","后轮增速 3.5：最大侧滑",7.1,2.35,TEAL); metric(s,"4.36°","后轮增速 0：最大侧滑",9.8,2.35,RED)
    text(s,"Interlagos A/B：0.133 s 漂移 → 0.0 s；当前策略的漂移 100% 来自后轮增速。",7.1,3.75,5.0,.45,16,INK,True)
    bullet(s,"这证实“在给定仿真执行器下”可做发卡弯漂移。",7.1,4.5,4.9)
    bullet(s,"不证实单电机四驱实车能复现：其后轮增速等效为 0。",7.1,5.1,4.9,RED); footer(s,9)

    s = new(); title(s, "ROBUSTNESS", "主动观测还不是鲁棒感知：组合扰动下 10/10 失败")
    image(s,"output/racing/mujoco/bahrain_c03v10c_sensor_perturb_seed10.png",.75,2.0,5.75,3.55)
    metric(s,"10/10","组合扰动失败",7.1,2.35,RED); metric(s,"2.7 s","最早撞墙时间",9.9,2.35,ORANGE)
    bullet(s,"条件：噪声 0.02 m + 丢束 5% + 延迟 50 ms；5 赛道 × 2 种子。",7.1,3.8,4.9)
    bullet(s,"消融：噪声单独可过；丢束或延迟单独已致命。",7.1,4.5,4.9)
    bullet(s,"根因：起步扫描匹配失败，进入 motion_estimate_unavailable 的盲车模式。",7.1,5.2,4.9,RED); footer(s,10)

    s = new(); title(s, "SIM-TO-REAL", "实车路线先恢复可观测性，再谈极限操控")
    steps=[("01","转向标定","舵机指令→左右轮实际角\n含 Ackermann 差动、滞后、速率限幅"),("02","补转向反馈","电位计/绝对编码器\n让安全前瞻使用实测转角"),("03","分级上线","shadow 回放→台架→低速→逐级提速\n每级留 rosbag/日志/圈速"),("04","漂移动作","接受不漂移，或新增独立后轮/手刹，或重训制动转移行为")]
    for i,(n,a,b) in enumerate(steps):
        x=.7+i*3.12; box(s,x,2.25,2.72,2.85,RGBColor(255,255,255),BLUE,True); text(s,n,x+.22,2.5,.35,.25,13,BLUE,True); text(s,a,x+.22,2.95,2.22,.35,17,INK,True); text(s,b,x+.22,3.55,2.22,.8,13,MUTED)
    text(s,"不要跳过阶段：两条“real”观测仿真初测有效，不等价于任何实车运行。",.7,5.75,11.5,.35,17,RED,True); footer(s,11)

    s = new(); title(s, "CLOSE", "可交付的是可复现证据；未闭合的问题也已被清楚定位")
    for x,a,b,c in [(0.8,"已完成","资产加载/渲染/短程步进\n双引擎录像与审计\nMu 24/24 整圈+超车",TEAL),(4.55,"部分完成","Isaac 22/24（7.25 m/s）\n9.0 m/s 速度榜\n留出种子证据",BLUE),(8.3,"下一步","扰动鲁棒扫描匹配\nIsaac 桥接触 / 起步接触\n转向反馈与漂移动作选择",ORANGE)]:
        box(s,x,2.2,3.3,2.65,RGBColor(255,255,255),c,True); text(s,a,x+.25,2.53,2.7,.3,22,c,True); text(s,b,x+.25,3.12,2.65,1.25,15,INK)
    text(s,"观看录像、复跑命令、结果口径：docs/RELEASE.md · docs/VALIDATION_STATUS.md · output/racing/media.html",.8,5.65,11.7,.38,15,NAVY,True)
    text(s,"谢谢。欢迎质疑结论，但请连同实验条件一起质疑。",.8,6.25,11.4,.35,18,INK,True); footer(s,12)
    OUT.parent.mkdir(parents=True, exist_ok=True); prs.save(OUT); print(OUT)

if __name__ == "__main__": build()
